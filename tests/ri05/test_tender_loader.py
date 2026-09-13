"""Unit tests for the tender-loading layer, against the committed Kessler Point package.

No network, no model, no key. The real tender in `data/tenders/kessler_point` is read
(never written); everything constructed points at pytest's tmp_path.

The counts below are pinned deliberately. They are what a re-read of the same folder must
keep producing, so a change to an extractor that quietly drops a page or a sheet fails
here rather than showing up as a bid that missed a requirement.
"""

import sys
from collections.abc import Iterator
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path

import pytest
from docx import Document as new_docx
from openpyxl import Workbook

from ri05_tender.tender.loader import (
    TenderLoadError,
    _sheet_lines,
    load_document,
    load_tender,
    main,
    render_table,
    sha256_of,
)
from ri05_tender.tender.models import DocumentPage, TenderDocument, TenderPackage, first_line_of

KESSLER_POINT = Path("data/tenders/kessler_point")

# Hand-checked against the committed package, and against both candidate PDF readers:
# 13 documents, 10 PDF (48 pages between them) and 3 XLSX (24 worksheets between them).
EXPECTED_DOCUMENTS = 13
EXPECTED_PDFS = 10
EXPECTED_XLSX = 3
EXPECTED_PDF_PAGES = 48
SHEET_COUNTS = {"07_Bill_of_Quantities": 10, "08_Pricing_Schedule": 12, "09_Compliance_Matrix": 2}


# Installed once for the whole session: an audit hook cannot be removed, so the cost of
# leaving it is one comparison per audit event, and it records only while _RECORDING is on.
_OPENED: list[str] = []
_RECORDING = False


def _audit(event: str, args: tuple[object, ...]) -> None:
    if _RECORDING and event == "open" and args and isinstance(args[0], str | bytes | Path):
        _OPENED.append(str(args[0]))


sys.addaudithook(_audit)


@contextmanager
def recording_opens() -> Iterator[list[str]]:
    """Collect every path opened inside the block."""
    global _RECORDING  # noqa: PLW0603 - an audit hook cannot be given per-call state
    _OPENED.clear()
    _RECORDING = True
    try:
        yield _OPENED
    finally:
        _RECORDING = False


@pytest.fixture(scope="module")
def kessler() -> TenderPackage:
    """The committed tender, loaded once for the whole module."""
    return load_tender(KESSLER_POINT)


def a_workbook(path: Path, sheets: dict[str, list[list[object]]]) -> Path:
    """Write a small .xlsx with named sheets and rows."""
    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(title=name)
        for row in rows:
            sheet.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path


def a_tender_folder(root: Path, *, with_gold: bool) -> Path:
    """Build a minimal but real tender folder, optionally with an answer key."""
    folder = root / "some_other_tender"
    (folder / "documents").mkdir(parents=True)
    a_workbook(folder / "documents" / "01_Schedule.xlsx", {"Sheet one": [["a", "b"], [1, 2]]})
    if with_gold:
        (folder / "gold").mkdir(parents=True)
        (folder / "gold" / "gold_set.jsonl").write_text('{"item_id": "x"}\n', encoding="utf-8")
    return folder


# --- the committed package ----------------------------------------------------------------


def test_the_committed_tender_loads_every_document(kessler: TenderPackage) -> None:
    assert len(kessler.documents) == EXPECTED_DOCUMENTS
    assert len(kessler.documents_of("pdf")) == EXPECTED_PDFS
    assert len(kessler.documents_of("xlsx")) == EXPECTED_XLSX
    assert kessler.documents_of("docx") == []


def test_the_pdfs_carry_every_printed_page(kessler: TenderPackage) -> None:
    total = sum(document.page_count for document in kessler.documents_of("pdf"))

    assert total == EXPECTED_PDF_PAGES


def test_the_committed_tender_has_a_gold_set(kessler: TenderPackage) -> None:
    assert kessler.has_gold
    assert kessler.gold_path == KESSLER_POINT / "gold" / "gold_set.jsonl"


def test_the_package_names_itself_from_the_folder(kessler: TenderPackage) -> None:
    # Not from a constant: the same code loads any tender folder.
    assert kessler.name == "kessler_point"
    assert kessler.root_path == KESSLER_POINT


def test_documents_are_ordered_and_identified_by_filename_stem(kessler: TenderPackage) -> None:
    assert kessler.document_ids[:2] == ["01_Instructions_to_Bidders", "02_Scope_of_Work"]
    assert kessler.document_ids == sorted(kessler.document_ids)


def test_reading_a_document_twice_gives_the_same_document() -> None:
    # Order and content are stable, which is what lets a sha256 mean anything. Checked on
    # one real PDF rather than by re-reading the whole package: extraction is where a
    # result could vary, and a second full read costs ten seconds to learn the same thing.
    path = KESSLER_POINT / "documents" / "04_Technical_Specification.pdf"

    assert load_document(path) == load_document(path)


# --- page anchoring: the assertion the whole layer exists for --------------------------------


def test_a_requirement_is_found_and_its_page_number_is_recorded(kessler: TenderPackage) -> None:
    # If this passes, an EvidenceRef can cite a page of the PDF the client actually sent.
    spec = kessler.document("04_Technical_Specification")
    assert spec is not None

    assert "TS-B.25" in spec.text
    pages = spec.pages_containing("TS-B.25")

    assert pages == [4]
    page = spec.page(4)
    assert page is not None
    assert "TS-B.25" in page.text
    assert page.document_id == "04_Technical_Specification"


def test_the_requirement_is_not_on_any_other_page(kessler: TenderPackage) -> None:
    # A page number is only worth citing if it is the page and not merely a page.
    spec = kessler.document("04_Technical_Specification")
    assert spec is not None

    elsewhere = [page.page_number for page in spec.pages if "TS-B.25" in page.text]

    assert elsewhere == [4]


def test_every_page_knows_which_document_it_belongs_to(kessler: TenderPackage) -> None:
    for document in kessler.documents:
        assert all(page.document_id == document.document_id for page in document.pages)


def test_every_document_is_numbered_from_one_without_gaps(kessler: TenderPackage) -> None:
    for document in kessler.documents:
        numbers = [page.page_number for page in document.pages]
        assert numbers == list(range(1, document.page_count + 1)), document.document_id


def test_pdf_text_is_prose_rather_than_tab_separated_words(kessler: TenderPackage) -> None:
    # This is the pdfplumber-over-pypdf choice, pinned. pypdf renders the same sentence as
    # "All\tcameras\tshall\tsupport", which cannot be quoted into an EvidenceRef.
    spec = kessler.document("04_Technical_Specification")
    assert spec is not None
    page = spec.page(4)
    assert page is not None

    assert "All cameras shall support" in page.text


# --- workbooks ------------------------------------------------------------------------------


@pytest.mark.parametrize(("document_id", "sheets"), sorted(SHEET_COUNTS.items()))
def test_each_workbook_yields_one_page_per_worksheet(
    kessler: TenderPackage, document_id: str, sheets: int
) -> None:
    document = kessler.document(document_id)
    assert document is not None

    assert len(document.sheet_names) == sheets
    assert document.page_count == sheets


def test_a_worksheet_page_starts_with_its_sheet_name(kessler: TenderPackage) -> None:
    matrix = kessler.document("09_Compliance_Matrix")
    assert matrix is not None

    for number, sheet_name in enumerate(matrix.sheet_names, start=1):
        page = matrix.page(number)
        assert page is not None
        assert page.text.splitlines()[0] == sheet_name


def test_a_workbooks_title_is_its_first_sheet_name(kessler: TenderPackage) -> None:
    matrix = kessler.document("09_Compliance_Matrix")
    assert matrix is not None

    assert matrix.title == matrix.sheet_names[0] == "Guidelines"


def test_only_workbooks_report_sheet_names(kessler: TenderPackage) -> None:
    assert all(document.sheet_names == [] for document in kessler.documents_of("pdf"))


def test_a_cell_yields_its_value_not_its_formula(tmp_path: Path) -> None:
    # data_only=True. A pricing schedule reporting "=SUM(B2:B40)" as its total would be
    # worse than useless, so this pins the read mode rather than trusting the default.
    path = a_workbook(tmp_path / "book.xlsx", {"Totals": [["item", "cost"], ["pump", 1250]]})

    document = load_document(path)

    assert "1250" in document.page(1).text  # type: ignore[union-attr]
    assert "=SUM" not in document.text


def test_a_row_keeps_its_row_number_so_a_row_citation_resolves(tmp_path: Path) -> None:
    # Line 1 is the sheet name, so line 1 + r is spreadsheet row r. A blank row in the
    # middle is kept for exactly this reason: dropping it would shift everything below.
    path = a_workbook(
        tmp_path / "boq.xlsx",
        {"A - PIDS": [["ref", "item"], [], ["1.2", "fence sensor"], ["1.3", "gate contact"]]},
    )

    lines = load_document(path).page(1).text.splitlines()  # type: ignore[union-attr]

    assert lines[0] == "A - PIDS"
    assert lines[1] == "ref\titem"
    assert lines[2] == ""  # row 2 was blank and is still row 2
    assert lines[3].startswith("1.2")
    assert lines[4].startswith("1.3")


def test_trailing_blank_rows_do_not_pad_the_page(tmp_path: Path) -> None:
    path = a_workbook(tmp_path / "padded.xlsx", {"S": [["a"], [], [], []]})

    assert load_document(path).page(1).text == "S\na"  # type: ignore[union-attr]


def test_the_trailing_trim_handles_a_sheet_that_over_reports_its_used_range() -> None:
    # Excel commonly declares a used range larger than the real content; a workbook
    # openpyxl wrote itself never does, so the round trip above cannot reach this and the
    # helper is exercised directly.
    rows = iter([("a", "b"), (None, None), ("c", None), (None, None), (None, None)])

    assert _sheet_lines(rows) == ["a\tb", "", "c"]


def test_a_date_cell_is_rendered_as_an_iso_string(tmp_path: Path) -> None:
    # So a deadline reads the same on every machine, whatever the locale. A bare date is
    # included because openpyxl hands it back as a datetime at midnight rather than a
    # date — checked against the installed version, and why the loader handles only one.
    path = a_workbook(
        tmp_path / "dates.xlsx",
        {"S": [["closes", datetime(2026, 9, 14, 17, 0)], ["issued", date(2026, 9, 1)]]},
    )

    text = load_document(path).page(1).text  # type: ignore[union-attr]

    assert "2026-09-14T17:00:00" in text
    assert "2026-09-01T00:00:00" in text


def test_an_empty_worksheet_is_still_a_page(tmp_path: Path) -> None:
    # Dropping it would renumber every sheet after it.
    path = a_workbook(tmp_path / "gap.xlsx", {"First": [["x"]], "Empty": [], "Third": [["y"]]})

    document = load_document(path)

    assert document.sheet_names == ["First", "Empty", "Third"]
    assert document.page(2).text == "Empty"  # type: ignore[union-attr]
    assert document.page(3).text.startswith("Third")  # type: ignore[union-attr]


# --- Word documents -------------------------------------------------------------------------


def a_word_file(path: Path) -> Path:
    """A .docx with a heading, a table and a paragraph after it."""
    document = new_docx()
    document.add_paragraph("Bidder Response Form")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Requirement"
    table.cell(0, 1).text = "Compliance"
    table.cell(1, 0).text = "TS-B.25"
    table.cell(1, 1).text = "Comply"
    document.add_paragraph("Signed on behalf of the bidder.")
    document.save(str(path))
    return path


def test_a_word_document_is_one_page_because_word_has_no_fixed_pagination(tmp_path: Path) -> None:
    document = load_document(a_word_file(tmp_path / "response.docx"))

    assert document.media_type == "docx"
    assert document.page_count == 1
    assert document.page(1).page_number == 1  # type: ignore[union-attr]


def test_word_table_text_is_read_and_not_silently_dropped(tmp_path: Path) -> None:
    # python-docx's `paragraphs` omits every table cell, and a compliance matrix is almost
    # entirely table. This is why the loader walks the body itself.
    document = load_document(a_word_file(tmp_path / "response.docx"))

    assert "TS-B.25" in document.text
    assert "Comply" in document.text


def test_word_content_keeps_its_document_order(tmp_path: Path) -> None:
    lines = load_document(a_word_file(tmp_path / "response.docx")).page(1).text.splitlines()  # type: ignore[union-attr]

    assert lines[0] == "Bidder Response Form"
    assert lines[-1] == "Signed on behalf of the bidder."


def test_a_word_documents_title_is_its_first_paragraph(tmp_path: Path) -> None:
    assert load_document(a_word_file(tmp_path / "r.docx")).title == "Bidder Response Form"


def test_a_word_document_reports_no_sheets(tmp_path: Path) -> None:
    assert load_document(a_word_file(tmp_path / "r.docx")).sheet_names == []


# --- the production path: a tender with no answer key ----------------------------------------


def test_a_tender_with_no_gold_folder_loads_identically(tmp_path: Path) -> None:
    # The production case, tested rather than assumed: a real tender arrives as documents
    # and nothing else.
    folder = a_tender_folder(tmp_path, with_gold=False)

    package = load_tender(folder)

    assert not package.has_gold
    assert package.gold_path is None
    assert package.document_ids == ["01_Schedule"]
    assert package.page_count == 1


def test_the_only_difference_a_gold_folder_makes_is_the_two_gold_fields(tmp_path: Path) -> None:
    without = load_tender(a_tender_folder(tmp_path / "a", with_gold=False))
    with_key = load_tender(a_tender_folder(tmp_path / "b", with_gold=True))

    assert without.model_dump(
        exclude={"has_gold", "gold_path", "root_path"}
    ) == with_key.model_dump(exclude={"has_gold", "gold_path", "root_path"})
    assert (without.has_gold, with_key.has_gold) == (False, True)


def test_a_gold_folder_without_the_gold_set_file_is_no_gold_set(tmp_path: Path) -> None:
    folder = a_tender_folder(tmp_path, with_gold=True)
    (folder / "gold" / "gold_set.jsonl").unlink()

    assert not load_tender(folder).has_gold


def test_the_loader_does_not_parse_the_gold_set(tmp_path: Path) -> None:
    # A loader that could read the answers is a loader that could leak them into the work.
    folder = a_tender_folder(tmp_path, with_gold=True)
    (folder / "gold" / "gold_set.jsonl").write_text("this is not JSON at all {{{", encoding="utf-8")

    package = load_tender(folder)

    assert package.has_gold
    assert package.gold_path is not None


def test_nothing_under_the_gold_folder_is_even_opened(tmp_path: Path) -> None:
    """Enforce data/tenders/README.md: no pipeline module reads anything under gold/.

    An audit hook rather than a patched `open`, because it sees every file opened by the
    process — including by pdfplumber, openpyxl and python-docx down in C — which a patched
    builtin would not.
    """
    folder = a_tender_folder(tmp_path, with_gold=True)
    gold = (folder / "gold").resolve()

    with recording_opens() as opened:
        package = load_tender(folder)

    assert package.has_gold  # it looked, so the negative below means something
    under_gold = [path for path in opened if gold in Path(path).resolve().parents]
    assert under_gold == []


# --- what the loader refuses -----------------------------------------------------------------


def test_a_missing_documents_folder_raises_naming_the_folder(tmp_path: Path) -> None:
    folder = tmp_path / "harbour_point"
    folder.mkdir()

    with pytest.raises(TenderLoadError) as error:
        load_tender(folder)

    assert "harbour_point" in str(error.value)
    assert "documents" in str(error.value)


def test_an_empty_documents_folder_raises_naming_the_folder(tmp_path: Path) -> None:
    folder = tmp_path / "harbour_point"
    (folder / "documents").mkdir(parents=True)

    with pytest.raises(TenderLoadError, match="harbour_point"):
        load_tender(folder)


def test_a_folder_that_is_not_a_directory_raises(tmp_path: Path) -> None:
    path = tmp_path / "not_a_folder.txt"
    path.write_text("x", encoding="utf-8")

    with pytest.raises(TenderLoadError, match="not a directory"):
        load_tender(path)


def test_an_unsupported_file_stops_the_load_rather_than_being_skipped(tmp_path: Path) -> None:
    # A document silently skipped is a document missing from the bid.
    folder = a_tender_folder(tmp_path, with_gold=False)
    (folder / "documents" / "02_Addendum.doc").write_text("old word format", encoding="utf-8")

    with pytest.raises(TenderLoadError, match="02_Addendum.doc"):
        load_tender(folder)


def test_a_directory_inside_documents_stops_the_load(tmp_path: Path) -> None:
    folder = a_tender_folder(tmp_path, with_gold=False)
    (folder / "documents" / "drawings").mkdir()

    with pytest.raises(TenderLoadError, match="drawings"):
        load_tender(folder)


def test_a_corrupt_document_names_the_file_and_the_format(tmp_path: Path) -> None:
    folder = a_tender_folder(tmp_path, with_gold=False)
    (folder / "documents" / "02_Broken.pdf").write_bytes(b"not a pdf at all")

    with pytest.raises(TenderLoadError) as error:
        load_tender(folder)

    assert "02_Broken.pdf" in str(error.value)
    assert "pdf" in str(error.value)


def test_dotfiles_are_ignored(tmp_path: Path) -> None:
    folder = a_tender_folder(tmp_path, with_gold=False)
    (folder / "documents" / ".DS_Store").write_bytes(b"\x00\x01")

    assert load_tender(folder).document_ids == ["01_Schedule"]


# --- the sha256 -------------------------------------------------------------------------------


def test_each_document_carries_the_hash_of_its_bytes(kessler: TenderPackage) -> None:
    document = kessler.document("13_Drawing_Register")
    assert document is not None

    assert document.sha256 == sha256_of(KESSLER_POINT / "documents" / document.filename)
    assert len(document.sha256) == 64


def test_a_changed_file_gets_a_different_hash(tmp_path: Path) -> None:
    folder = a_tender_folder(tmp_path, with_gold=False)
    before = load_tender(folder).documents[0].sha256
    a_workbook(folder / "documents" / "01_Schedule.xlsx", {"Sheet one": [["a", "b"], [1, 999]]})

    assert load_tender(folder).documents[0].sha256 != before


def test_hashes_are_unique_across_the_committed_package(kessler: TenderPackage) -> None:
    hashes = [document.sha256 for document in kessler.documents]

    assert len(set(hashes)) == len(hashes)


# --- the models' own guarantees -----------------------------------------------------------------


def a_document(**overrides: object) -> TenderDocument:
    settings: dict[str, object] = {
        "document_id": "01_Doc",
        "filename": "01_Doc.pdf",
        "media_type": "pdf",
        "title": "A document",
        "pages": [DocumentPage(document_id="01_Doc", page_number=1, text="one")],
        "sha256": "a" * 64,
    }
    settings.update(overrides)
    return TenderDocument.model_validate(settings)


def test_a_page_filed_under_another_documents_id_is_refused() -> None:
    with pytest.raises(ValueError, match="belonging to"):
        a_document(pages=[DocumentPage(document_id="99_Elsewhere", page_number=1, text="x")])


def test_page_numbers_with_a_gap_are_refused() -> None:
    # "page 7" has to be unambiguous or there was no point carrying page numbers.
    with pytest.raises(ValueError, match="without gaps"):
        a_document(
            pages=[
                DocumentPage(document_id="01_Doc", page_number=1, text="a"),
                DocumentPage(document_id="01_Doc", page_number=3, text="b"),
            ]
        )


def test_page_numbers_starting_at_zero_are_refused() -> None:
    with pytest.raises(ValueError):
        DocumentPage(document_id="01_Doc", page_number=0, text="a")


def test_sheet_names_on_something_that_is_not_a_workbook_are_refused() -> None:
    with pytest.raises(ValueError, match="carries sheet names"):
        a_document(sheet_names=["Sheet1"])


def test_a_workbook_whose_sheets_and_pages_disagree_is_refused() -> None:
    with pytest.raises(ValueError, match="one page per sheet"):
        a_document(media_type="xlsx", filename="01_Doc.xlsx", sheet_names=["A", "B"])


def test_two_documents_with_one_id_are_refused() -> None:
    with pytest.raises(ValueError, match="duplicate document id"):
        TenderPackage(
            name="t",
            root_path=Path("/tmp/t"),
            documents=[a_document(), a_document(filename="01_Doc.xlsx")],
            has_gold=False,
        )


def test_a_package_cannot_claim_a_gold_set_it_cannot_point_at() -> None:
    with pytest.raises(ValueError, match="gold_path is None"):
        TenderPackage(name="t", root_path=Path("/tmp/t"), has_gold=True)


def test_a_package_cannot_point_at_a_gold_set_it_denies_having() -> None:
    with pytest.raises(ValueError, match="has_gold is False"):
        TenderPackage(
            name="t", root_path=Path("/tmp/t"), has_gold=False, gold_path=Path("/tmp/t/gold.jsonl")
        )


def test_a_loaded_document_cannot_be_edited(kessler: TenderPackage) -> None:
    with pytest.raises(ValueError):
        kessler.documents[0].title = "something else"  # type: ignore[misc]


def test_searching_a_document_can_ignore_case(kessler: TenderPackage) -> None:
    spec = kessler.document("04_Technical_Specification")
    assert spec is not None

    assert spec.pages_containing("ts-b.25") == []
    assert spec.pages_containing("ts-b.25", case_sensitive=False) == [4]


def test_asking_for_a_page_that_does_not_exist_gives_none(kessler: TenderPackage) -> None:
    spec = kessler.document("04_Technical_Specification")
    assert spec is not None

    assert spec.page(9999) is None
    assert kessler.document("no_such_document") is None


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("", ""),
        ("\n\n  \n", ""),
        ("  Document  4  —  Spec \nnext line", "Document 4 — Spec"),
        ("first\nsecond", "first"),
    ],
)
def test_a_title_is_the_first_non_empty_line_with_whitespace_collapsed(
    text: str, expected: str
) -> None:
    assert first_line_of(text) == expected


def test_a_title_is_capped_so_a_runaway_first_line_cannot_fill_the_record() -> None:
    assert len(first_line_of("x" * 500)) == 200


# --- the command line ---------------------------------------------------------------------------


def test_the_cli_loads_a_folder_and_prints_a_row_per_document(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    # Wiring only, against a small folder. What the table says about the real tender is
    # asserted through `render_table` below, off the already-loaded fixture, because a
    # third full read of 48 real PDF pages costs ten seconds to prove the same thing.
    folder = a_tender_folder(tmp_path, with_gold=True)

    exit_code = main([str(folder)])

    out = capsys.readouterr().out
    assert exit_code == 0
    assert "document_id" in out and "sha256" in out
    assert "01_Schedule" in out
    assert "1 document(s)" in out
    assert "Gold set: yes" in out


def test_the_cli_says_when_there_is_no_gold_set(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    main([str(a_tender_folder(tmp_path, with_gold=False))])

    assert "Gold set: no" in capsys.readouterr().out


def test_the_cli_reports_a_bad_folder_and_exits_non_zero(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    exit_code = main([str(tmp_path / "nowhere")])

    assert exit_code == 1
    assert "Could not load the tender" in capsys.readouterr().out


def test_the_table_shows_every_document_with_its_type_pages_and_hash(
    kessler: TenderPackage,
) -> None:
    table = render_table(kessler)

    for document in kessler.documents:
        assert document.document_id in table
        assert document.sha256 in table
    assert "xlsx" in table and "pdf" in table
    assert f"{EXPECTED_DOCUMENTS} document(s)" in table
    assert f"{EXPECTED_PDF_PAGES + sum(SHEET_COUNTS.values())} page(s)" in table
    assert f"{sum(SHEET_COUNTS.values())} worksheet(s)" in table


def test_the_gold_folder_guard_can_actually_fail(tmp_path: Path) -> None:
    # Guards the guard. A watcher that never sees anything would pass the test above
    # whether or not the loader behaved, which is no test at all.
    folder = a_tender_folder(tmp_path, with_gold=True)
    gold = (folder / "gold").resolve()

    with recording_opens() as opened:
        (folder / "gold" / "gold_set.jsonl").read_text(encoding="utf-8")

    under_gold = [path for path in opened if gold in Path(path).resolve().parents]
    assert under_gold != []
