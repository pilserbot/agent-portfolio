"""Reading a tender folder off disk into typed documents and pages.

`load_tender(folder)` takes the tender folder itself — a `Path`, not a name and not a
constant — and returns a `TenderPackage`. Every tender enters the system through this one
function, and no tender's name appears anywhere in this module. The layout is the contract::

    <tender folder>/
        documents/   the tender as the client issued it
        gold/        an answer key. Optional, and absent in production.

Invoked as a command for a first look at a new folder::

    python -m ri05_tender.tender.loader data/tenders/<folder>

**A folder with no `gold/` loads identically.** That is the production path: a real tender
arrives as documents and nothing else. `has_gold` is set by looking for
`gold/gold_set.jsonl` and the file is never opened here — a loader that could read the
answers is a loader that could leak them into the work.

## Choosing the PDF reader

pdfplumber, not pypdf, and the two were run against the committed tender before choosing
rather than picked from memory. They agree on page counts (48 across the ten PDFs) and both
find `TS-B.25` on page 4 of the Technical Specification, so page anchoring is safe either
way. The text is not:

- pdfplumber gives ``TS-B.5 All cameras shall support HTTPS with TLS 1.3, SRTP, IEEE
  802.1X certificate-based authentication, …`` — one sentence, as printed.
- pypdf gives the same sentence with a tab between every single word, and a line break
  wherever the styling changes, so the bolded phrase "HTTPS with TLS 1.3" is split onto
  its own line and the comma after it onto another.

That text cannot be quoted into an `EvidenceRef` and would corrupt any requirement
extracted from it. pdfplumber costs about 2.4× the time (10.1s against 4.2s for all ten
documents) and is worth it: this runs once per tender, and the text it produces is the
input to everything after.

## What the extraction does and does not promise

- **PDF** — one page per printed page, in order. What the client paginated is what a
  citation names.
- **XLSX** — one page per worksheet, numbered from 1 in book order, read with
  ``data_only=True`` so a cell yields its value and not its formula. The text is the sheet
  name on the first line, then one line per row, so **line 1 + r is spreadsheet row r**.
  Blank rows in the middle of a sheet are kept for exactly that reason: dropping them would
  shift every row number below, and "BOQ row 214" has to still mean row 214.
- **DOCX** — the whole document as page 1. Word has no fixed pagination: where a page
  breaks depends on the renderer, the fonts installed and the paper size, so any page
  number this module invented would be a different number on the next machine. One page is
  the honest answer.

  Tables are read in document order alongside the paragraphs, because `python-docx`'s
  `paragraphs` silently omits every table cell and a compliance matrix is almost entirely
  table. Two further traps, both found by building a file that contains them rather than by
  reading the docs: a **merged cell** is yielded once per column it spans and would repeat
  its text across the line, and a **table nested inside a cell** is invisible to
  `cell.text` in the same way a top-level table is invisible to `paragraphs`. Both are
  handled, and both are pinned by a test — see `_table_lines`.

Deliberately does not: call a model, reach a network, interpret anything it reads, or open
the gold set. It identifies no requirement, parses no clause and judges no answer — it
turns files into text with the page numbers kept. It also does not repair a document it
cannot read: an unreadable or unsupported file in `documents/` stops the load, because a
document silently skipped is a document missing from the bid.
"""

import argparse
import hashlib
from collections.abc import Iterator, Sequence
from datetime import datetime
from pathlib import Path, PurePosixPath

import pdfplumber
from docx import Document as open_docx
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

from ri05_tender.tender.models import (
    DocumentPage,
    MediaType,
    TenderDocument,
    TenderPackage,
    first_line_of,
)

DOCUMENTS_DIRNAME = "documents"

# The one place in this package that may name the answer key, and it may only be tested for
# existence. `tests/ri05/test_no_gold_leakage.py` parses every module here and fails on any
# other reference to it; this single literal is the whole of its allowance. Kept as one
# constant rather than a directory and a filename so that allowance is one line, and so
# widening it cannot happen quietly.
GOLD_SET_RELATIVE_PATH = PurePosixPath("gold/gold_set.jsonl")

# The suffixes this loader can read. Anything else in documents/ stops the load rather than
# being skipped: the set is closed because adding a format means teaching this module to
# read it, not relaxing a filter.
SUFFIX_TO_MEDIA_TYPE: dict[str, MediaType] = {
    ".pdf": "pdf",
    ".xlsx": "xlsx",
    ".docx": "docx",
}

_HASH_CHUNK_BYTES = 1 << 20

__all__ = [
    "DOCUMENTS_DIRNAME",
    "GOLD_SET_RELATIVE_PATH",
    "SUFFIX_TO_MEDIA_TYPE",
    "TenderLoadError",
    "load_document",
    "load_tender",
    "main",
    "sha256_of",
]


class TenderLoadError(Exception):
    """A tender folder could not be read as a tender.

    Always names the path, because the first thing anyone asks of this error is which
    folder it is talking about.
    """


def sha256_of(path: Path) -> str:
    """The sha256 of a file's bytes, read in chunks so a large drawing set does not load."""
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(_HASH_CHUNK_BYTES):
            digest.update(chunk)
    return digest.hexdigest()


def _cell_text(value: object) -> str:
    """Render one spreadsheet cell as text without inventing precision.

    Dates become ISO strings so a deadline reads the same on every machine and in every
    locale; everything else is rendered as Python already renders it. An empty cell is an
    empty string, which is what keeps the row-per-line alignment intact.

    Only `datetime` is handled, not `date`: openpyxl returns a date-formatted cell as a
    `datetime` at midnight, verified against the installed version, so a branch for bare
    dates would never run.
    """
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value)


def _pdf_pages(path: Path, document_id: str) -> tuple[list[DocumentPage], str]:
    """Every printed page of a PDF, in order, with the first line of page 1 as the title."""
    pages: list[DocumentPage] = []
    with pdfplumber.open(path) as pdf:
        for number, page in enumerate(pdf.pages, start=1):
            pages.append(
                DocumentPage(
                    document_id=document_id,
                    page_number=number,
                    text=page.extract_text() or "",
                )
            )
    return pages, first_line_of(pages[0].text) if pages else ""


def _sheet_lines(rows: Iterator[tuple[object, ...]]) -> list[str]:
    """One line per spreadsheet row, with trailing blank rows trimmed.

    Interior blanks are kept so line numbers keep tracking row numbers; only the run of
    empty rows at the end is dropped, because a sheet saved by Excel commonly declares a
    used range larger than its real content and would otherwise end in hundreds of blanks.
    A workbook openpyxl wrote itself never has that problem, so the trim is covered by a
    direct test rather than by a round trip through a file.
    """
    lines = ["\t".join(_cell_text(cell) for cell in row).rstrip("\t") for row in rows]
    while lines and not lines[-1].strip():
        lines.pop()
    return lines


def _xlsx_pages(path: Path, document_id: str) -> tuple[list[DocumentPage], list[str], str]:
    """One page per worksheet, numbered from 1 in book order.

    `data_only=True` reads what a cell evaluated to rather than the formula behind it: a
    pricing schedule that reported "=SUM(B2:B40)" as its total would be worse than useless.
    """
    pages: list[DocumentPage] = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_names = list(workbook.sheetnames)
        for number, sheet_name in enumerate(sheet_names, start=1):
            worksheet = workbook[sheet_name]
            body = _sheet_lines(worksheet.iter_rows(values_only=True))
            pages.append(
                DocumentPage(
                    document_id=document_id,
                    page_number=number,
                    text="\n".join([sheet_name, *body]),
                )
            )
    finally:
        workbook.close()
    return pages, sheet_names, sheet_names[0] if sheet_names else ""


def _cell_content(cell: _Cell) -> tuple[str, list[str]]:
    """One cell's own text, and the lines of any table nested inside it.

    A cell's paragraphs are joined with a space so the cell stays one column of its row;
    a nested table cannot, so it comes back separately to be emitted as its own rows.
    """
    own: list[str] = []
    nested: list[str] = []
    for item in cell.iter_inner_content():
        if isinstance(item, Paragraph):
            if item.text:
                own.append(item.text)
        else:
            nested.extend(_table_lines(item))
    return " ".join(own), nested


def _table_lines(table: Table) -> list[str]:
    """One line per table row, cells tab-separated, nested tables following their row.

    Two things Word does that a naive walk gets wrong, both found by probing a file with
    them in it rather than by reading the docs:

    - A merged cell is yielded by `row.cells` once for every column it spans, so its text
      would appear two or three times across one line. `row.cells` returns the *same* cell
      object each time, so identity is what de-duplicates it.
    - A table nested inside a cell is invisible to `cell.text`, exactly as a top-level
      table is invisible to `document.paragraphs`. Its rows are emitted after the row that
      holds it, so nothing is dropped.
    """
    lines: list[str] = []
    for row in table.rows:
        texts: list[str] = []
        nested: list[str] = []
        seen: list[_Cell] = []
        for cell in row.cells:
            if any(cell is already for already in seen):
                continue
            seen.append(cell)
            text, inner = _cell_content(cell)
            texts.append(text)
            nested.extend(inner)
        lines.append("\t".join(texts).rstrip("\t"))
        lines.extend(nested)
    return lines


def _docx_blocks(path: Path) -> list[str]:
    """Paragraph and table text in document order.

    Through `iter_inner_content()`, which is public and yields paragraphs and tables in the
    order they appear. `document.paragraphs` omits every table cell — verified against the
    installed python-docx before relying on it — and a compliance matrix is almost entirely
    table, so reading only the paragraphs would lose the document.
    """
    document = open_docx(str(path))
    blocks: list[str] = []
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            blocks.append(item.text)
        else:
            blocks.extend(_table_lines(item))
    return blocks


def _docx_pages(path: Path, document_id: str) -> tuple[list[DocumentPage], str]:
    """The whole Word document as page 1. See the module docstring for why there is only one."""
    blocks = _docx_blocks(path)
    text = "\n".join(blocks)
    page = DocumentPage(document_id=document_id, page_number=1, text=text)
    return [page], first_line_of(text)


def load_document(path: Path) -> TenderDocument:
    """Read one tender document into pages.

    The document id is the filename stem, so a citation survives the file being moved
    between folders but not renamed — which is the right sensitivity, since a renamed
    tender document is a different document as far as anyone reading the bid is concerned.
    """
    media_type = SUFFIX_TO_MEDIA_TYPE.get(path.suffix.lower())
    if media_type is None:
        raise TenderLoadError(
            f"{path}: cannot read a {path.suffix or '(no suffix)'} file. Supported: "
            f"{', '.join(sorted(SUFFIX_TO_MEDIA_TYPE))}."
        )

    document_id = path.stem
    sheet_names: list[str] = []
    try:
        if media_type == "pdf":
            pages, title = _pdf_pages(path, document_id)
        elif media_type == "xlsx":
            pages, sheet_names, title = _xlsx_pages(path, document_id)
        else:
            pages, title = _docx_pages(path, document_id)
    except Exception as error:  # noqa: BLE001 - three libraries, each with its own failures
        raise TenderLoadError(f"{path}: could not be read as {media_type}: {error}") from error

    return TenderDocument(
        document_id=document_id,
        filename=path.name,
        media_type=media_type,
        title=title,
        pages=pages,
        sheet_names=sheet_names,
        sha256=sha256_of(path),
    )


def _document_paths(documents_dir: Path, folder: Path) -> list[Path]:
    """Every file in documents/, sorted by name, dotfiles ignored.

    Sorted so two loads of one folder produce the same package in the same order. A
    directory inside documents/ is an error rather than something to walk into or skip
    quietly: the layout is flat, and a folder of drawings nobody noticed would be a folder
    of drawings nobody bid on.
    """
    entries = [entry for entry in documents_dir.iterdir() if not entry.name.startswith(".")]
    directories = sorted(entry.name for entry in entries if entry.is_dir())
    if directories:
        noun = "directory" if len(directories) == 1 else "directories"
        raise TenderLoadError(
            f"{folder}: {DOCUMENTS_DIRNAME}/ holds {noun} ({', '.join(directories)}); the "
            f"tender layout is flat, so every document must sit directly in "
            f"{DOCUMENTS_DIRNAME}/."
        )
    return sorted((entry for entry in entries if entry.is_file()), key=lambda path: path.name)


def load_tender(folder: Path) -> TenderPackage:
    """Read a tender folder into a `TenderPackage`.

    Takes the folder itself. Raises `TenderLoadError`, naming the folder, when it is not a
    directory, when `documents/` is missing or empty, or when a file in it cannot be read.

    A folder with no `gold/` loads exactly as one with it; only `has_gold` and `gold_path`
    differ. Nothing under `gold/` is opened.
    """
    folder = Path(folder)
    if not folder.is_dir():
        raise TenderLoadError(f"{folder}: not a directory, so it cannot be a tender folder.")

    documents_dir = folder / DOCUMENTS_DIRNAME
    if not documents_dir.is_dir():
        raise TenderLoadError(
            f"{folder}: has no {DOCUMENTS_DIRNAME}/ folder. A tender is a folder containing "
            f"{DOCUMENTS_DIRNAME}/ (the files the client issued) and optionally "
            f"{GOLD_SET_RELATIVE_PATH.parent}/ (an answer key)."
        )

    paths = _document_paths(documents_dir, folder)
    if not paths:
        raise TenderLoadError(
            f"{folder}: {DOCUMENTS_DIRNAME}/ is empty. There is nothing to bid on."
        )

    gold_file = folder / GOLD_SET_RELATIVE_PATH
    has_gold = gold_file.is_file()

    return TenderPackage(
        name=folder.name,
        root_path=folder,
        documents=[load_document(path) for path in paths],
        has_gold=has_gold,
        gold_path=gold_file if has_gold else None,
    )


def render_table(package: TenderPackage) -> str:
    """The package as a table, for a first look at a folder nobody has loaded before."""
    header = ("document_id", "type", "pages", "sha256")
    rows = [
        (
            document.document_id,
            document.media_type,
            str(document.page_count),
            document.sha256,
        )
        for document in package.documents
    ]
    widths = [max(len(str(cell)) for cell in column) for column in zip(header, *rows, strict=True)]

    def line(cells: Sequence[str]) -> str:
        left = "  ".join(
            cell.ljust(width) for cell, width in zip(cells[:2], widths[:2], strict=True)
        )
        pages = cells[2].rjust(widths[2])
        return f"{left}  {pages}  {cells[3]}"

    gold = f"yes ({package.gold_path})" if package.has_gold else "no"
    sheets = sum(len(document.sheet_names) for document in package.documents)
    return "\n".join(
        [
            f"Tender: {package.name}   ({package.root_path})",
            f"Gold set: {gold}",
            "",
            line(header),
            "  ".join("-" * width for width in widths),
            *(line(row) for row in rows),
            "",
            f"{len(package.documents)} document(s), {package.page_count} page(s), "
            f"{sheets} worksheet(s).",
        ]
    )


def main(argv: Sequence[str] | None = None) -> int:
    """Run the command-line entry point and return the process exit code."""
    parser = argparse.ArgumentParser(
        prog="python -m ri05_tender.tender.loader",
        description="Read a tender folder and print what it holds.",
    )
    parser.add_argument(
        "folder",
        type=Path,
        help="The tender folder: the one containing documents/ and optionally gold/.",
    )
    args = parser.parse_args(argv)

    try:
        package = load_tender(args.folder)
    except TenderLoadError as error:
        print(f"Could not load the tender: {error}")
        return 1

    print(render_table(package))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
